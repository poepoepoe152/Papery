"""Document workload sweep: formats, big docs, hostile inputs — via real API."""
import io
import time
import uuid

import requests

BASE = "http://127.0.0.1:8000/api/v1"
results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond)))
    print(f"[{'PASS' if cond else 'FAIL'}] {name}" + (f"  <- {detail}" if not cond else (f"  ({detail})" if detail else "")))


def api(method, path, token=None, **kw):
    headers = kw.pop("headers", {})
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return requests.request(method, f"{BASE}{path}", headers=headers, timeout=120, **kw)


ROWS = [("Booking No", "SITGUAYE12345"), ("Container No", "TCLU1234567"),
        ("Seal No", "ML-998877"), ("Gross Weight", "26,557.68 KG"),
        ("Consignee", "GLOBEX LOGISTICS"), ("Port of Loading", "LAEM CHABANG")]


def pdf_bytes(rows, title="BOOKING CONFIRMATION", pages=1):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for p in range(pages):
        c.setFont("Helvetica-Bold", 15); c.drawString(50, 800, title if p == 0 else f"PAGE {p+1} RIDER")
        c.setFont("Helvetica", 11)
        y = 760
        if p == 0:
            for k, v in rows:
                c.drawString(50, y, f"{k}: {v}"); y -= 24
        else:
            for i in range(30):
                c.drawString(50, y, f"Line item {p}-{i}: CARGO DESCRIPTION CONTINUED"); y -= 24
        c.showPage()
    c.save()
    return buf.getvalue()


def docx_bytes(rows, title="SHIPPING INSTRUCTION"):
    import docx
    d = docx.Document()
    d.add_heading(title, 0)
    for k, v in rows:
        d.add_paragraph(f"{k}: {v}")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def png_bytes(rows, title="DRAFT BILL OF LADING"):
    """Render a form PDF page to PNG (goes through the OCR path)."""
    import fitz
    data = pdf_bytes(rows, title)
    doc = fitz.open(stream=data, filetype="pdf")
    pix = doc[0].get_pixmap(dpi=200)
    return pix.tobytes("png")


def tiff_bytes(rows, title="PACKING LIST"):
    from PIL import Image
    import fitz
    data = pdf_bytes(rows, title)
    doc = fitz.open(stream=data, filetype="pdf")
    pix = doc[0].get_pixmap(dpi=150)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    buf = io.BytesIO(); img.save(buf, "TIFF")
    return buf.getvalue()


# --- setup: fresh company on Pro plan ---
rid = uuid.uuid4().hex[:6]
r = api("POST", "/auth/register", json={"company_name": f"DocCo-{rid}", "full_name": "Doc Tester",
                                        "email": f"doc-{rid}@t.co", "password": "Password123"})
T = r.json()["access_token"]
# generate a fresh Pro license via super admin
sa = api("POST", "/auth/login", json={"email": "admin@papery.app", "password": "Admin123!"}).json()["access_token"]
plans = api("GET", "/admin/plans", token=sa).json()
pro = next(p for p in plans if p["name"] == "Pro")
key = api("POST", "/admin/licenses", token=sa, json={"plan_id": pro["id"]}).json()["license_key"]
r = api("POST", "/license/activate", token=T, json={"license_key": key})
check("setup: fresh Pro company", r.status_code == 200, key)


def upload(name, data, mime="application/octet-stream", expect=201):
    r = api("POST", "/files/upload", token=T, files={"file": (name, data, mime)})
    return r


def run_verification(title, ref_ids, tgt_ids, timeout=180):
    r = api("POST", "/verifications", token=T,
            json={"title": title, "reference_file_ids": ref_ids, "target_file_ids": tgt_ids})
    if r.status_code != 201:
        return None, f"create {r.status_code} {r.text[:80]}", 0
    vid = r.json()["id"]
    t0 = time.time()
    while time.time() - t0 < timeout:
        st = api("GET", f"/verifications/{vid}/status", token=T).json()
        if st["status"] in ("COMPLETED", "FAILED"):
            return vid, st, round(time.time() - t0, 1)
        time.sleep(1.5)
    return vid, {"status": "TIMEOUT"}, round(time.time() - t0, 1)


ref_pdf = upload("ref.pdf", pdf_bytes(ROWS), "application/pdf").json()["id"]

# --- 1) DOCX target ---
tgt_rows = [(k, v) for k, v in ROWS]
tgt_rows[1] = ("Container No", "TCLU1234561")  # deliberate error
f = upload("si.docx", docx_bytes(tgt_rows), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
check("upload DOCX", f.status_code == 201, str(f.status_code))
vid, st, dt = run_verification("pdf vs docx", [ref_pdf], [f.json()["id"]])
det = api("GET", f"/verifications/{vid}", token=T).json() if vid else {}
check("DOCX verification", st.get("status") == "COMPLETED" and det.get("critical_count") == 1,
      f"{st.get('status')} crit={det.get('critical_count')} in {dt}s")

# --- 2) PNG image target (OCR path) ---
f = upload("photo.png", png_bytes(tgt_rows), "image/png")
check("upload PNG", f.status_code == 201, str(f.status_code))
vid, st, dt = run_verification("pdf vs png(ocr)", [ref_pdf], [f.json()["id"]])
det = api("GET", f"/verifications/{vid}", token=T).json() if vid else {}
ocr_crit = det.get("critical_count", -1)
check("PNG (OCR) verification completes", st.get("status") == "COMPLETED",
      f"{st.get('status')} in {dt}s crit={ocr_crit}")
check("PNG (OCR) catches container error", ocr_crit is not None and ocr_crit >= 1, f"crit={ocr_crit}")

# --- 3) TIFF target ---
f = upload("scan.tiff", tiff_bytes(tgt_rows), "image/tiff")
check("upload TIFF", f.status_code == 201, str(f.status_code))
vid, st, dt = run_verification("pdf vs tiff(ocr)", [ref_pdf], [f.json()["id"]])
check("TIFF verification completes", st.get("status") == "COMPLETED", f"{st.get('status')} in {dt}s")

# --- 4) multi-page (5p) target ---
f = upload("tgt5p.pdf", pdf_bytes(tgt_rows, "DRAFT BILL OF LADING", pages=5), "application/pdf")
vid, st, dt = run_verification("multipage 5p", [ref_pdf], [f.json()["id"]])
det = api("GET", f"/verifications/{vid}", token=T).json() if vid else {}
check("5-page verification", st.get("status") == "COMPLETED" and det.get("critical_count") == 1,
      f"{st.get('status')} crit={det.get('critical_count')} in {dt}s")

# --- 5) 300-page stress ---
big = pdf_bytes(tgt_rows, "DRAFT BILL OF LADING", pages=300)
print(f"  ... 300-page pdf size = {len(big)//1024} KB")
f = upload("big300.pdf", big, "application/pdf")
check("upload 300p PDF", f.status_code == 201, str(f.status_code))
vid, st, dt = run_verification("stress 300 pages", [ref_pdf], [f.json()["id"]], timeout=300)
check("300-page verification", st.get("status") == "COMPLETED", f"{st.get('status')} in {dt}s")
print(f"  ... 300-page processing time: {dt}s")

# --- 6) corrupted PDF ---
f = upload("corrupt.pdf", b"%PDF-1.7 then total garbage \x00\xff\x13" * 100, "application/pdf")
check("upload corrupt PDF accepted", f.status_code == 201, str(f.status_code))
vid, st, dt = run_verification("corrupt target", [ref_pdf], [f.json()["id"]])
det = api("GET", f"/verifications/{vid}", token=T).json() if vid else {}
print(f"  ... corrupt doc result: status={st.get('status')} overall={det.get('overall_result')} "
      f"findings={len(det.get('findings', []))}")
check("corrupt doc does NOT silently PASS",
      not (st.get("status") == "COMPLETED" and det.get("overall_result") == "PASS"),
      f"status={st.get('status')} overall={det.get('overall_result')}")

# --- 7) PNG bytes disguised as .pdf ---
f = upload("fake.pdf", png_bytes(tgt_rows), "application/pdf")
vid, st, dt = run_verification("fake ext", [ref_pdf], [f.json()["id"]])
det = api("GET", f"/verifications/{vid}", token=T).json() if vid else {}
check("disguised PNG-as-PDF does NOT silently PASS",
      not (st.get("status") == "COMPLETED" and det.get("overall_result") == "PASS"),
      f"status={st.get('status')} overall={det.get('overall_result')}")

# --- 8) oversized file ---
r = upload("huge.pdf", b"%PDF" + b"0" * (26 * 1024 * 1024), "application/pdf")
check("oversized 26MB -> 413", r.status_code == 413, str(r.status_code))

fails = [n for n, ok in results if not ok]
print(f"\n===== {len(results)} checks | PASS {len(results)-len(fails)} | FAIL {len(fails)} =====")
for n in fails:
    print("  FAILED:", n)
