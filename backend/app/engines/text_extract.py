"""Text extraction with automatic OCR, orientation correction, and preprocessing.

- Digital PDFs: extract the text layer directly (PyMuPDF).
- Scanned PDFs / images: OCR via Tesseract (PaddleOCR is the production engine;
  Tesseract is the bundled fallback from the architecture's OCR chain).
- Phone photos are often rotated; standalone images get automatic orientation
  detection (try 0/90/180/270, keep the most text-like result) plus grayscale +
  autocontrast, which dramatically improves OCR on real-world captures.
- DOCX: read paragraphs and tables.

OCR degrades gracefully: if the Tesseract binary is unavailable, extraction
still returns whatever digital text exists rather than raising.
"""
import io
import os
import re

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"}
_FINAL_MAX_SIDE = 2400
_THUMB_MAX_SIDE = 1400


def _word_score(text: str) -> int:
    return len(re.findall(r"[A-Za-z]{3,}", text))


def _ocr_raw(image, config: str = "") -> str:
    try:
        import pytesseract

        return pytesseract.image_to_string(image, config=config)
    except Exception as exc:  # noqa: BLE001 - missing binary or runtime error
        print(f"[ocr] skipped: {exc}")
        return ""


def _detect_orientation(image) -> int:
    """Return the best rotation (degrees) by trying all four on a thumbnail."""
    try:
        from PIL import ImageOps

        thumb = image.copy()
        thumb.thumbnail((_THUMB_MAX_SIDE, _THUMB_MAX_SIDE))
        thumb = ImageOps.grayscale(thumb)
    except Exception:  # noqa: BLE001
        return 0

    best_angle, best_score = 0, -1
    for angle in (0, 90, 180, 270):
        rotated = thumb.rotate(-angle, expand=True)
        score = _word_score(_ocr_raw(rotated, config="--psm 6"))
        if score > best_score:
            best_angle, best_score = angle, score
    return best_angle


def _ocr_image(image, detect_orientation: bool = True) -> str:
    from PIL import ImageOps

    image = ImageOps.exif_transpose(image)
    if detect_orientation:
        angle = _detect_orientation(image)
        if angle:
            image = image.rotate(-angle, expand=True)

    work = image.copy()
    work.thumbnail((_FINAL_MAX_SIDE, _FINAL_MAX_SIDE))
    work = ImageOps.autocontrast(ImageOps.grayscale(work))
    return _ocr_raw(work)


def _extract_pdf(path: str) -> tuple[str, int, list]:
    import fitz  # PyMuPDF

    parts: list[str] = []
    words: list = []  # (page_index, x0, y0, x1, y1, text) for layout-aware extraction
    with fitz.open(path) as doc:
        page_count = doc.page_count
        for page_index, page in enumerate(doc):
            text = page.get_text("text")
            if len(text.strip()) < 20:
                # Likely a scanned page -> rasterize and OCR (orientation from
                # a PDF render is usually already upright, so skip the 4x scan).
                try:
                    from PIL import Image

                    pix = page.get_pixmap(dpi=200)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = _ocr_image(img, detect_orientation=False)
                except Exception as exc:  # noqa: BLE001
                    print(f"[ocr] page render failed: {exc}")
            else:
                for w in page.get_text("words"):
                    words.append((page_index, w[0], w[1], w[2], w[3], w[4]))
            parts.append(text)
    return "\n".join(parts), page_count, words


def _extract_docx(path: str) -> tuple[str, int]:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts), 1


def _extract_image(path: str) -> tuple[str, int]:
    from PIL import Image

    with Image.open(path) as img:
        return _ocr_image(img, detect_orientation=True), 1


def extract_text(path: str, original_name: str = "") -> dict:
    """Return {text, page_count, engine, words}.

    `words` holds per-word coordinates for digital PDFs (empty otherwise) and
    feeds the layout-aware extractor for multi-column forms.
    """
    ext = os.path.splitext(original_name or path)[1].lower()
    words: list = []
    try:
        if ext == ".pdf":
            text, pages, words = _extract_pdf(path)
            engine = "PDF/OCR"
        elif ext == ".docx":
            text, pages = _extract_docx(path)
            engine = "DOCX"
        elif ext in _IMAGE_EXTS:
            text, pages = _extract_image(path)
            engine = "OCR"
        else:
            text, pages, engine = "", 0, "UNSUPPORTED"
    except Exception as exc:  # noqa: BLE001
        print(f"[extract] failed for {original_name}: {exc}")
        text, pages, engine = "", 0, "ERROR"

    return {"text": text or "", "page_count": pages, "engine": engine, "words": words}
