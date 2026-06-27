"""Verification report generation: JSON, Excel (XLSX), and PDF."""
import io
import json
from datetime import datetime, timezone


def _findings_payload(findings) -> list[dict]:
    return [
        {
            "field": f.field_label or f.field_key,
            "reference_value": f.reference_value,
            "detected_value": f.detected_value,
            "match_type": f.match_type,
            "severity": f.severity,
            "explanation": f.explanation,
            "suggested_fix": f.suggested_fix,
            "status": f.status,
        }
        for f in findings
    ]


def build_json(verification, documents, findings) -> bytes:
    payload = {
        "verification": {
            "id": str(verification.id),
            "title": verification.title,
            "status": verification.status,
            "overall_result": verification.overall_result,
            "critical_count": verification.critical_count,
            "major_count": verification.major_count,
            "minor_count": verification.minor_count,
            "generated_at": datetime.now(timezone.utc).isoformat(),
        },
        "documents": [
            {
                "role": d.role,
                "name": d.original_name,
                "doc_type": d.doc_type,
                "fields": d.fields or {},
            }
            for d in documents
        ],
        "findings": _findings_payload(findings),
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def build_xlsx(verification, documents, findings) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"

    ws.append(["Papery Verification Report"])
    ws["A1"].font = Font(size=14, bold=True)
    ws.append(["Title", verification.title])
    ws.append(["Result", verification.overall_result])
    ws.append([
        "Critical / Major / Minor",
        f"{verification.critical_count} / {verification.major_count} / {verification.minor_count}",
    ])
    ws.append([])

    header = ["Field", "Reference Value", "Detected Value", "Severity",
              "Match Type", "Explanation", "Suggested Fix"]
    ws.append(header)
    header_row = ws.max_row
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")

    colors = {"CRITICAL": "FECACA", "MAJOR": "FEF3C7", "MINOR": "E5E7EB"}
    for f in findings:
        ws.append([
            f.field_label or f.field_key, f.reference_value, f.detected_value,
            f.severity, f.match_type, f.explanation, f.suggested_fix,
        ])
        fill = colors.get(f.severity)
        if fill:
            ws.cell(row=ws.max_row, column=4).fill = PatternFill("solid", fgColor=fill)

    for col, width in zip("ABCDEFG", [22, 28, 28, 12, 14, 40, 28]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(verification, documents, findings) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Papery Verification Report")
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Papery Verification Report", styles["Title"]))
    elements.append(Paragraph(verification.title or "", styles["Heading2"]))
    result_color = {"PASS": "green", "WARN": "orange", "FAIL": "red"}.get(
        verification.overall_result or "", "black"
    )
    elements.append(
        Paragraph(
            f"Result: <font color='{result_color}'><b>{verification.overall_result}</b></font>"
            f" &nbsp; Critical: {verification.critical_count} &nbsp;"
            f" Major: {verification.major_count} &nbsp; Minor: {verification.minor_count}",
            styles["Normal"],
        )
    )
    elements.append(Spacer(1, 8 * mm))

    data = [["Field", "Reference", "Detected", "Severity", "Suggested Fix"]]
    for f in findings:
        data.append([
            Paragraph(str(f.field_label or f.field_key), styles["BodyText"]),
            Paragraph(str(f.reference_value or "—"), styles["BodyText"]),
            Paragraph(str(f.detected_value or "—"), styles["BodyText"]),
            f.severity,
            Paragraph(str(f.suggested_fix or "—"), styles["BodyText"]),
        ])

    table = Table(data, colWidths=[32 * mm, 38 * mm, 38 * mm, 20 * mm, 38 * mm], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]
    sev_color = {"CRITICAL": "#FECACA", "MAJOR": "#FEF3C7", "MINOR": "#E5E7EB"}
    for i, f in enumerate(findings, start=1):
        c = sev_color.get(f.severity)
        if c:
            style.append(("BACKGROUND", (3, i), (3, i), colors.HexColor(c)))
    table.setStyle(TableStyle(style))
    elements.append(table)

    if not findings:
        elements.append(Paragraph("No differences detected.", styles["Normal"]))

    doc.build(elements)
    return buf.getvalue()


def build_report(fmt: str, verification, documents, findings) -> tuple[bytes, str, str]:
    """Return (bytes, file_extension, mime_type)."""
    fmt = fmt.upper()
    if fmt == "JSON":
        return build_json(verification, documents, findings), "json", "application/json"
    if fmt == "XLSX":
        return (
            build_xlsx(verification, documents, findings),
            "xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if fmt == "PDF":
        return build_pdf(verification, documents, findings), "pdf", "application/pdf"
    raise ValueError(f"Unsupported report format: {fmt}")
